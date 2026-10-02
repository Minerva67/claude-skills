#!/usr/bin/env python3
"""Verify cleaning is pure deletion, scan for residue, optionally write back to a
spreadsheet column, and build a side-by-side HTML review page.

Pairs files by name: <orig-dir>/row<N>.txt  <->  <clean-dir>/row<N>.txt

The core guarantee: cleaned text must be obtainable from the original by DELETION ONLY.
A difflib pass flags any insert/replace (i.e. characters the cleaner added or changed) as
a VIOLATION — those rows must be fixed before delivery, because they break the "fair
variant of the original" promise.

Usage:
  python verify_and_report.py --orig-dir work/orig --clean-dir work/clean --html work/compare.html
  # optional write-back into a new column of the source sheet:
  python verify_and_report.py --orig-dir work/orig --clean-dir work/clean --html work/compare.html \
      --xlsx source.xlsx --col G --out-xlsx cleaned.xlsx --new-col-title "清洗后prompt"
"""
import argparse, os, re, html, difflib, glob

RESIDUE = [
    ('bracket-tag', re.compile(r'\[(?:CODE|Code-gen|DIFF|Diffusion|HYB|Hybrid|3D|2D|CG|COMP)[^\]]*\]')),
    ('paradigm-word', re.compile(r'\b(?:diffusion|code-?gen|hybrid|code-drawn)\b', re.I)),
    ('engine/tech-stack', re.compile(r'\b(?:Remotion|React-Three-Fiber|R3F|Three\.?js|GSAP|D3|WebGL|KaTeX|Lottie|Blender|Houdini|Flux|SDXL|Runway|Sora|Kling|Veo|Pixi)\b')),
    ('tech-section', re.compile(r'(?:Technical Notes|Render notes|Code-gen stack|Diffusion approach|Compositing law|render pipeline|render method)', re.I)),
]

def classify(text):
    tl = text.lower()
    if re.search(r'\[(?:code|diff|hyb|3d|2d|cg|comp)', tl): return 'Bracket render tag — names which paradigm/technique renders this shot.'
    if re.search(r'remotion|three\.?js|r3f|react-three|gsap|\bd3\b|webgl|katex|lottie|blender|houdini|flux|sdxl|runway|sora|kling|\bveo\b|pixi', tl): return 'Engine / tech-stack / model name — a specific implementation tool.'
    if re.search(r'diffusion|code-?gen|hybrid|code-drawn', tl): return 'Paradigm / model word (diffusion / code-gen / hybrid) — biases toward a model.'
    if re.search(r'technical notes|production note|render notes|code-gen stack|diffusion approach', tl): return 'Whole technical-implementation passage (stack / render notes).'
    if re.search(r'compositing law|composite:|overlay|\bplate\b', tl): return 'Hybrid compositing / pipeline description (plate + overlay).'
    if re.search(r'stack:|diffusion usage|render method|render pipeline|paste .{0,40}tool', tl): return 'Tech-stack / tooling declaration.'
    if re.search(r'^\s*\*\*.{0,40}:\*\*|base \[|overlay \[|^\s*3d:|^\s*2d:', tl): return 'Per-shot technical prefix / render-split label (label removed, creative text kept).'
    return 'Technical-implementation content.'

def esc(s): return html.escape(s)
def esc_a(s): return html.escape(s, quote=True)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--orig-dir', required=True)
    ap.add_argument('--clean-dir', required=True)
    ap.add_argument('--html', required=True)
    ap.add_argument('--xlsx'); ap.add_argument('--col'); ap.add_argument('--out-xlsx')
    ap.add_argument('--sheet', default=None)
    ap.add_argument('--new-col-title', default='cleaned prompt')
    a = ap.parse_args()

    pairs = []
    for op in sorted(glob.glob(os.path.join(a.orig_dir, 'row*.txt')),
                     key=lambda p: int(re.search(r'row(\d+)', p).group(1))):
        n = int(re.search(r'row(\d+)', op).group(1))
        cp = os.path.join(a.clean_dir, f'row{n}.txt')
        if os.path.exists(cp):
            pairs.append((n, op, cp))

    blocks = []; total_del = total_bad = 0; violations = []; residue = {}
    for n, op, cp in pairs:
        orig = open(op).read(); clean = open(cp).read()
        o = orig.rstrip('\n'); c = clean.rstrip('\n')
        sm = difflib.SequenceMatcher(None, o, c, autojunk=False)
        ndel = nins = nrep = 0; parts = []; di = 0
        for tag, i1, i2, j1, j2 in sm.get_opcodes():
            if tag == 'equal': parts.append(esc(o[i1:i2]))
            elif tag == 'delete':
                ndel += 1; di += 1; num = f'row{n}-{di}'; seg = o[i1:i2]
                parts.append(f'<span class="del" data-id="{num}" data-rs="{esc_a(classify(seg))}">{esc(seg)}<sup>{num}</sup></span>')
            elif tag == 'insert':
                nins += 1; parts.append(f'<span class="ins" title="ADDED by cleaner (violation)">{esc(c[j1:j2])}</span>')
            elif tag == 'replace':
                nrep += 1; parts.append(f'<span class="rep" title="CHANGED by cleaner (violation): -> {esc_a(c[j1:j2])}">{esc(o[i1:i2])}</span>')
        bad = nins + nrep; total_del += ndel; total_bad += bad
        if bad: violations.append((n, nins, nrep))
        # residue scan
        rhits = []
        for name, pat in RESIDUE:
            f = pat.findall(c)
            if f: rhits.append(f'{name}: ' + ', '.join(dict.fromkeys(f))[:200])
        if rhits: residue[n] = rhits
        warn = f' <span class="bad">⚠ {bad} violation(s)</span>' if bad else ''
        res = f' <span class="res">⚠ residue</span>' if rhits else ''
        blocks.append(f'''<section id="r{n}"><h2>row{n} <span class="cnt">{ndel} deleted</span>{warn}{res}</h2>
<div class="grid"><div class="col l"><div class="tag">original</div><pre>{esc(orig)}</pre></div>
<div class="col r"><div class="tag">cleaned (red = deleted; hover for reason)</div><pre>{''.join(parts)}</pre></div></div></section>''')

    # write-back
    if a.xlsx and a.col and a.out_xlsx:
        import openpyxl
        from openpyxl.utils import get_column_letter
        from openpyxl.styles import Font, Alignment, PatternFill
        wb = openpyxl.load_workbook(a.xlsx, data_only=False)
        ws = wb[a.sheet] if a.sheet else wb.active
        newc = ws.max_column + 1; L = get_column_letter(newc)
        h = ws.cell(row=1, column=newc, value=a.new_col_title)
        h.font = Font(bold=True, color='FFFFFF'); h.fill = PatternFill('solid', fgColor='305496')
        h.alignment = Alignment(horizontal='center', vertical='center')
        for n, op, cp in pairs:
            c = open(cp).read().rstrip('\n')
            cell = ws.cell(row=n, column=newc, value=c)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[L].width = 85
        wb.save(a.out_xlsx)
        print(f'Wrote cleaned prompts to column {L} of {a.out_xlsx}')

    nav = ' '.join(f'<a href="#r{n}">row{n}</a>' for n, _, _ in pairs)
    body = '\n'.join(blocks)
    page = f'''<!DOCTYPE html><html lang="zh"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Prompt cleaning review</title>
<style>*{{box-sizing:border-box}}body{{margin:0;font-family:-apple-system,"PingFang SC","Microsoft YaHei",sans-serif;background:#f4f5f7;color:#1a1a1a}}
header{{position:sticky;top:0;z-index:5;background:#fff;border-bottom:1px solid #e2e4e8;padding:14px 22px}}header h1{{margin:0 0 6px;font-size:16px}}.sub{{font-size:13px;color:#666;line-height:1.6}}
.navbar{{padding:8px 22px;background:#fbfbfc;border-bottom:1px solid #eee;font-size:13px;line-height:2}}.navbar a{{color:#2563c9;text-decoration:none;margin-right:6px}}
section{{padding:14px 22px;border-bottom:8px solid #eceef1}}section h2{{margin:0 0 10px;font-size:15px}}.cnt{{font-size:12px;color:#888;font-weight:400;margin-left:6px}}.bad{{font-size:12px;color:#fff;background:#c0392b;padding:1px 8px;border-radius:10px;margin-left:6px}}.res{{font-size:12px;color:#7a5c00;background:#fff2cc;padding:1px 8px;border-radius:10px;margin-left:6px}}
.grid{{display:grid;grid-template-columns:1fr 1fr;gap:0}}.col{{padding:12px 16px;min-width:0}}.col.r{{border-left:1px solid #e2e4e8;background:#fcfcfd}}.tag{{font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:#999;margin-bottom:8px}}
pre{{margin:0;white-space:pre-wrap;word-break:break-word;font-family:"SF Mono",ui-monospace,Menlo,Consolas,monospace;font-size:12.5px;line-height:1.7}}
.del{{color:#c0392b;text-decoration:line-through;background:#fdeaea;border-radius:3px;padding:0 1px;cursor:help}}.del sup{{font-size:9px;color:#8a1f14;text-decoration:none;margin-left:1px}}
.ins{{background:#8e44ad;color:#fff;border-radius:3px;padding:0 2px}}.rep{{background:#f39c12;color:#000;border-radius:3px;padding:0 2px}}
#tt{{position:fixed;max-width:380px;background:#1f2430;color:#f5f6f8;padding:11px 13px;border-radius:9px;box-shadow:0 8px 26px rgba(0,0,0,.28);font-size:13px;line-height:1.5;z-index:50;pointer-events:none;opacity:0;transition:opacity .1s}}#tt .id{{font-size:11px;color:#e0a800;margin-bottom:5px}}
</style></head><body>
<header><h1>Prompt cleaning review — {len(pairs)} prompts · {total_del} spans deleted · {total_bad} violations</h1>
<div class="sub"><b style="color:#c0392b">red strike</b> = deleted technical content (numbered, hover for why); <b style="color:#8e44ad">purple</b>/<b style="color:#f39c12">orange</b> = cleaner added/changed text (violations, must fix). Kept text is byte-for-byte identical to the original.</div></header>
<div class="navbar">{nav}</div>{body}<div id="tt"></div>
<script>const tt=document.getElementById('tt');document.querySelectorAll('.del').forEach(el=>{{
el.addEventListener('mouseenter',()=>{{tt.innerHTML='<div class="id">'+el.dataset.id+'</div><div>'+el.dataset.rs+'</div>';tt.style.opacity='1';}});
el.addEventListener('mousemove',e=>{{let x=e.clientX+15,y=e.clientY+15;const w=tt.offsetWidth,h=tt.offsetHeight;if(x+w>innerWidth)x=e.clientX-w-15;if(y+h>innerHeight)y=e.clientY-h-15;tt.style.left=x+'px';tt.style.top=y+'px';}});
el.addEventListener('mouseleave',()=>tt.style.opacity='0');}});</script></body></html>'''
    open(a.html, 'w', encoding='utf-8').write(page)

    print(f'\nPrompts: {len(pairs)} | spans deleted: {total_del} | violations: {total_bad}')
    if violations:
        print('VIOLATIONS (only-delete broken — fix these rows):')
        for n, i, r in violations: print(f'  row{n}: {i} inserted, {r} replaced')
    else:
        print('Only-delete check PASSED (0 violations) — kept text is byte-for-byte original.')
    if residue:
        print('Residue (verify each — real miss vs. false positive):')
        for n, hits in residue.items():
            print(f'  row{n}: ' + ' | '.join(hits))
    else:
        print('Residue scan: clean.')
    print(f'Review page: {a.html}')

if __name__ == '__main__':
    main()
