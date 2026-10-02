#!/usr/bin/env python3
"""Export each prompt cell from a spreadsheet column to its own text file.

Writes <out-dir>/row<N>.txt for every non-empty cell (N = spreadsheet row number),
so downstream cleaning and verification operate on plain text with stable ids.

Usage:
  python extract_column.py --xlsx file.xlsx --col G --out-dir work/orig [--id-col A] [--sheet Sheet1] [--start-row 2]
  python extract_column.py --csv file.csv --col 6 --out-dir work/orig [--id-col 0]   # 0-based for CSV
"""
import argparse, os, sys, csv

def col_to_idx(col):
    col = col.strip()
    if col.isdigit():
        return int(col)  # caller-defined convention; used directly for CSV (0-based)
    idx = 0
    for ch in col.upper():
        idx = idx * 26 + (ord(ch) - ord('A') + 1)
    return idx  # 1-based for openpyxl

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx'); ap.add_argument('--csv')
    ap.add_argument('--col', required=True, help='xlsx: column letter (G) or 1-based num; csv: 0-based num')
    ap.add_argument('--id-col', default=None, help='optional id/key column (same convention as --col)')
    ap.add_argument('--sheet', default=None)
    ap.add_argument('--out-dir', required=True)
    ap.add_argument('--start-row', type=int, default=2)
    a = ap.parse_args()
    os.makedirs(a.out_dir, exist_ok=True)
    exported = []

    if a.xlsx:
        import openpyxl
        wb = openpyxl.load_workbook(a.xlsx, data_only=False)
        ws = wb[a.sheet] if a.sheet else wb.active
        ci = col_to_idx(a.col); idi = col_to_idx(a.id_col) if a.id_col else None
        for r in range(a.start_row, ws.max_row + 1):
            v = ws.cell(row=r, column=ci).value
            if not (isinstance(v, str) and v.strip()):
                continue
            open(os.path.join(a.out_dir, f'row{r}.txt'), 'w').write(v)
            key = ws.cell(row=r, column=idi).value if idi else ''
            exported.append((r, key))
    elif a.csv:
        ci = int(a.col); idi = int(a.id_col) if a.id_col is not None else None
        with open(a.csv, newline='') as f:
            rows = list(csv.reader(f))
        for r in range(a.start_row - 1, len(rows)):
            row = rows[r]
            if ci >= len(row): continue
            v = row[ci]
            if not (isinstance(v, str) and v.strip()): continue
            open(os.path.join(a.out_dir, f'row{r+1}.txt'), 'w').write(v)
            key = row[idi] if (idi is not None and idi < len(row)) else ''
            exported.append((r + 1, key))
    else:
        sys.exit('Provide --xlsx or --csv')

    print(f'Exported {len(exported)} prompts to {a.out_dir}')
    print('row -> id:')
    for r, k in exported:
        print(f'  row{r}\t{k}')

if __name__ == '__main__':
    main()
